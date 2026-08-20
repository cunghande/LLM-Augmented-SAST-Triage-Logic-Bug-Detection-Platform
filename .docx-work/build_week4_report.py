"""Build the professional Week 4 SAST triage progress report."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path(__file__).with_name("Bao_cao_tien_do_Tuan_4_SAST_Triage.docx")
FONT = "Times New Roman"
BLACK = RGBColor(0, 0, 0)
GRAY = RGBColor(89, 89, 89)
LIGHT_GRAY = "E7E6E6"
VERY_LIGHT_GRAY = "F7F7F7"
BORDER_GRAY = "BFBFBF"
TABLE_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120


def set_run_font(
    run,
    *,
    size: float = 13,
    bold: bool = False,
    italic: bool = False,
    color: RGBColor = BLACK,
) -> None:
    """Apply the required Times New Roman typography to a run."""
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{key}"), FONT)


def set_style_font(style, *, size: float, bold: bool = False) -> None:
    """Set a Word style with deterministic Latin and East Asian font mappings."""
    style.font.name = FONT
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = BLACK
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{key}"), FONT)


def configure_styles(doc: Document) -> None:
    """Configure the report style system with the user-requested typography."""
    normal = doc.styles["Normal"]
    set_style_font(normal, size=13)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.3

    for name, before, after in (
        ("Heading 1", 14, 6),
        ("Heading 2", 10, 5),
        ("Heading 3", 8, 4),
    ):
        style = doc.styles[name]
        set_style_font(style, size=14, bold=True)
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.15
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True

    title = doc.styles["Title"]
    set_style_font(title, size=14, bold=True)
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(8)
    title.paragraph_format.line_spacing = 1.15
    title.paragraph_format.keep_with_next = True

    subtitle = doc.styles["Subtitle"]
    set_style_font(subtitle, size=13)
    subtitle.font.italic = False
    subtitle.font.color.rgb = BLACK
    subtitle.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_before = Pt(0)
    subtitle.paragraph_format.space_after = Pt(4)
    subtitle.paragraph_format.line_spacing = 1.15

    caption = doc.styles["Caption"]
    set_style_font(caption, size=13, bold=True)
    caption.font.italic = False
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(4)
    caption.paragraph_format.keep_with_next = True


def set_cell_shading(cell, fill: str) -> None:
    """Apply a grayscale cell fill."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, *, top: int = 100, bottom: int = 100, start: int = 120, end: int = 120) -> None:
    """Set table cell margins in DXA."""
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("bottom", bottom), ("start", start), ("end", end)):
        element = tc_mar.find(qn(f"w:{tag}"))
        if element is None:
            element = OxmlElement(f"w:{tag}")
            tc_mar.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa: int) -> None:
    """Set a cell's exact width in DXA."""
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa: list[int]) -> None:
    """Apply fixed, internally consistent Word table geometry."""
    if sum(widths_dxa) != TABLE_WIDTH_DXA:
        raise ValueError("Table widths must total 9360 DXA")

    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(TABLE_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")

    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), BORDER_GRAY)

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            set_cell_width(cell, widths_dxa[index])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def repeat_table_header(row) -> None:
    """Mark a table row as a repeating header row."""
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def format_cell(cell, *, bold: bool = False, align=WD_ALIGN_PARAGRAPH.LEFT) -> None:
    """Format all text in one table cell."""
    for paragraph in cell.paragraphs:
        paragraph.alignment = align
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.15
        for run in paragraph.runs:
            set_run_font(run, size=13, bold=bold)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths_dxa: list[int]):
    """Add a fixed-width grayscale report table."""
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for index, text in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = text
        set_cell_shading(cell, LIGHT_GRAY)
        format_cell(cell, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    repeat_table_header(table.rows[0])

    for row_values in rows:
        cells = table.add_row().cells
        for index, text in enumerate(row_values):
            cells[index].text = text
            if len(table.rows) % 2 == 1:
                set_cell_shading(cells[index], VERY_LIGHT_GRAY)
            format_cell(cells[index])
    set_table_geometry(table, widths_dxa)

    return table


def add_heading(
    doc: Document,
    text: str,
    level: int = 1,
    *,
    page_break_before: bool = False,
) -> None:
    """Add a black, bold, 14-point section heading."""
    paragraph = doc.add_paragraph(text, style=f"Heading {level}")
    paragraph.paragraph_format.page_break_before = page_break_before
    for run in paragraph.runs:
        set_run_font(run, size=14, bold=True)


def add_body(doc: Document, text: str, *, keep_with_next: bool = False) -> None:
    """Add a justified 13-point body paragraph."""
    paragraph = doc.add_paragraph(style="Normal")
    paragraph.paragraph_format.keep_with_next = keep_with_next
    run = paragraph.add_run(text)
    set_run_font(run, size=13)


def add_label_paragraph(doc: Document, label: str, text: str) -> None:
    """Add a body paragraph with a restrained bold lead label."""
    paragraph = doc.add_paragraph(style="Normal")
    label_run = paragraph.add_run(label)
    set_run_font(label_run, size=13, bold=True)
    text_run = paragraph.add_run(text)
    set_run_font(text_run, size=13)


def add_page_field(paragraph) -> None:
    """Append a live PAGE field to a footer paragraph."""
    run = paragraph.add_run()
    set_run_font(run, size=10, color=GRAY)
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.extend((fld_char_begin, instr_text, fld_char_end))


def configure_page(doc: Document) -> None:
    """Configure the standard business brief page geometry and page furniture."""
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hp.paragraph_format.space_after = Pt(0)
    hr = hp.add_run("LLM-Augmented SAST Triage | Báo cáo Tuần 4")
    set_run_font(hr, size=10, color=GRAY)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.paragraph_format.space_before = Pt(0)
    fp.paragraph_format.space_after = Pt(0)
    fr = fp.add_run("Báo cáo tiến độ Tuần 4 | Trang ")
    set_run_font(fr, size=10, color=GRAY)
    add_page_field(fp)


def build_report() -> Path:
    """Create the Week 4 report and return its path."""
    doc = Document()
    configure_styles(doc)
    configure_page(doc)

    doc.core_properties.title = "Báo cáo tiến độ Tuần 4 - LLM-Augmented SAST Triage"
    doc.core_properties.subject = "Tiến độ dự án và trọng tâm phát hiện lỗi logic"
    doc.core_properties.author = "Nhóm phát triển dự án"
    doc.core_properties.keywords = "SAST, LLM, logic bug, IDOR, mass assignment"

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(14)

    title = doc.add_paragraph("BÁO CÁO TIẾN ĐỘ TUẦN 4", style="Title")
    for run in title.runs:
        set_run_font(run, size=14, bold=True)
    subtitle = doc.add_paragraph("DỰ ÁN LLM-AUGMENTED SAST TRIAGE", style="Subtitle")
    for run in subtitle.runs:
        set_run_font(run, size=13, bold=True)
    meta = doc.add_paragraph("Kỳ báo cáo: Tuần 4 - cập nhật đến ngày 17/08/2026", style="Subtitle")
    for run in meta.runs:
        set_run_font(run, size=13)
    status = doc.add_paragraph(
        "Trạng thái: Đang triển khai và kiểm chứng chức năng phát hiện lỗi logic",
        style="Subtitle",
    )
    status.paragraph_format.space_after = Pt(16)
    for run in status.runs:
        set_run_font(run, size=13, bold=True)

    add_heading(doc, "TÓM TẮT TIẾN ĐỘ", 1)
    add_body(
        doc,
        "Dự án đã hoàn thành nền tảng CLI triage cho repository Python: chạy Semgrep và "
        "Gitleaks, chuẩn hóa finding, lấy ngữ cảnh an toàn, làm sạch dữ liệu nhạy cảm, gọi "
        "LLM, áp dụng policy bảo thủ và xuất báo cáo JSON/Markdown. Hai chế độ triage "
        "artifact và scan-and-triage sử dụng chung một pipeline; mọi finding vẫn được giữ "
        "trong báo cáo kể cả khi provider gặp lỗi hoặc bằng chứng chưa đủ.",
    )
    add_body(
        doc,
        "Trong Tuần 4, dự án chuyển trọng tâm sang phát hiện lỗi logic mà các rule SAST "
        "thông thường khó nhận diện đầy đủ. Để bảo đảm khối lượng phù hợp và có thể kiểm "
        "chứng sâu, phạm vi được chủ động giới hạn ở hai nhóm ưu tiên: thiếu kiểm soát phân "
        "quyền/IDOR và mass assignment. Các nhóm lỗi logic khác chưa được đưa vào mục tiêu "
        "nghiệm thu của giai đoạn này.",
    )

    add_table(
        doc,
        ["Hạng mục", "Trạng thái hiện tại", "Kết quả chính"],
        [
            ["Pipeline triage cốt lõi", "Đã hoàn thành", "CLI end-to-end, một pipeline dùng chung, xử lý lỗi theo từng finding"],
            ["An toàn dữ liệu và ngữ cảnh", "Đã hoàn thành phần lõi", "Redact secret, kiểm tra path containment, sanitize context trước LLM"],
            ["Evidence và repository index", "Đã vận hành", "AST Python, evidence same-file/cross-file có giới hạn và metadata source-free"],
            ["Tích hợp provider LLM", "Đã tích hợp, đang hiệu chỉnh", "Fake provider cho test; Modal/Qwen và Ollama cho thử nghiệm thực"],
            ["Phát hiện lỗi logic", "Đang thực hiện", "Đã tạo candidate và evidence; đang tập trung kiểm chứng hai nhóm ưu tiên"],
        ],
        [2400, 2160, 4800],
    )

    add_heading(doc, "PHẦN ĐÃ HOÀN THÀNH", 1)
    add_label_paragraph(
        doc,
        "Pipeline và mô hình miền. ",
        "Các contract public được tách rõ, dùng Pydantic frozen model và bảo đảm raw source "
        "không vượt ranh giới component. Policy giữ vai trò quyết định cuối cùng thay vì "
        "để LLM tự động loại bỏ finding.",
    )
    add_label_paragraph(
        doc,
        "Scanner và báo cáo. ",
        "Semgrep/Gitleaks đã được đưa vào luồng chạy tuần tự, có timeout, artifact nội bộ "
        "và report JSON/Markdown. Reporter thực hiện kiểm tra rò rỉ phòng thủ trước khi ghi file.",
    )
    add_label_paragraph(
        doc,
        "Ngữ cảnh và evidence. ",
        "SafeContextService kiểm tra traversal, absolute path và symlink escape; context được "
        "redact/sanitize trước khi trở thành TriageInput. Repository index hỗ trợ theo dõi "
        "symbol, call-site và luồng source-guard-sink trong phạm vi có giới hạn.",
    )
    add_label_paragraph(
        doc,
        "Kiểm thử kỹ thuật. ",
        "Lần xác minh ngày 17/08/2026 đạt 479 test passed, 2 skipped; Ruff báo All checks "
        "passed trên src, tests, modal_apps và scripts.",
    )

    add_heading(doc, "TRỌNG TÂM TUẦN 4: PHÁT HIỆN LỖI LOGIC", 1)
    add_body(
        doc,
        "Detector sử dụng AST Python và repository index để tìm luồng dữ liệu, guard bảo vệ "
        "và thao tác nhạy cảm. Kết quả ở giai đoạn này là candidate có bằng chứng kèm theo, "
        "không phải kết luận lỗ hổng tự động. Candidate tiếp tục đi qua Safe Context, LLM, "
        "schema validation và policy bảo thủ; trường hợp chưa đủ bằng chứng được chuyển sang "
        "MANUAL_REVIEW.",
    )

    add_heading(doc, "Thiếu kiểm soát phân quyền / IDOR", 2)
    add_label_paragraph(
        doc,
        "Lý do ưu tiên. ",
        "Đây là nhóm lỗi có thể cho phép người dùng đọc, sửa hoặc xóa đối tượng không thuộc "
        "quyền sở hữu chỉ bằng cách thay đổi định danh trong request. Ảnh hưởng có thể bao "
        "gồm rò rỉ dữ liệu, thay đổi dữ liệu trái phép và vượt quyền nghiệp vụ.",
    )
    add_label_paragraph(
        doc,
        "Cách phát hiện đang triển khai. ",
        "Hệ thống theo dõi object identifier từ route, query hoặc body đến thao tác load, "
        "update hay delete; sau đó kiểm tra có guard phân quyền đúng đối tượng, truy vấn đã "
        "scope theo user/tenant hoặc dependency authorization tương ứng hay không. Bộ kiểm "
        "tra counter-evidence giúp tránh cảnh báo khi đã có guard hợp lệ.",
    )
    add_label_paragraph(
        doc,
        "Trạng thái hiện tại. ",
        "Candidate, evidence card và đường đi qua pipeline đã hoạt động. Trên artifact "
        "benchmark v5, phạm vi IDOR có 24 case đã review: 9 true positive, 0 false positive, "
        "1 false negative và 14 true negative; F1 = 0,9474. Đây là kết quả của detector xác "
        "định trên các location đã gán nhãn, chưa phải chất lượng tổng quát của LLM hoặc toàn repository.",
    )

    add_heading(doc, "Mass assignment", 2)
    add_label_paragraph(
        doc,
        "Lý do ưu tiên. ",
        "Lỗi xuất hiện khi dữ liệu do người dùng kiểm soát được gán hàng loạt vào model hoặc "
        "đối tượng nghiệp vụ mà không có allowlist trường được phép. Kẻ tấn công có thể chèn "
        "các thuộc tính nhạy cảm như role, owner, trạng thái phê duyệt hoặc hạn mức.",
    )
    add_label_paragraph(
        doc,
        "Cách phát hiện đang triển khai. ",
        "Detector tìm các luồng từ request payload đến constructor, update hoặc setattr hàng "
        "loạt; đồng thời kiểm tra allowlist, schema DTO, exclude field và phép gán từng trường "
        "để loại trừ trường hợp đã kiểm soát.",
    )
    add_label_paragraph(
        doc,
        "Trạng thái hiện tại. ",
        "Bộ phân tích đã sinh candidate và remediation/checklist theo lớp lỗi. Trên artifact "
        "benchmark v5 có 20 case dương đã review, detector nhận đúng 18 case và bỏ sót 2 case. "
        "Dataset này chưa có case âm được review cho mass assignment, vì vậy chưa sử dụng "
        "precision hoặc false-positive rate để tuyên bố chất lượng tổng quát.",
    )

    doc.add_page_break()

    add_heading(doc, "KẾT QUẢ KỸ THUẬT HIỆN TẠI", 1)
    add_table(
        doc,
        ["Nội dung xác minh", "Kết quả", "Cách hiểu"],
        [
            ["Full regression test", "479 passed / 2 skipped", "Toàn bộ test suite hiện chạy xanh"],
            ["Static quality", "Ruff: All checks passed", "Không phát hiện lỗi lint trong phạm vi kiểm tra"],
            ["IDOR - benchmark đã review", "9 TP / 0 FP / 1 FN / 14 TN", "Có cả mẫu dương và âm; F1 = 0,9474"],
            ["Mass assignment - benchmark đã review", "18 TP / 2 FN", "Chỉ có mẫu dương; chưa đủ cơ sở kết luận precision"],
            ["An toàn pipeline", "Fail closed", "Provider lỗi hoặc evidence thiếu không làm mất finding"],
        ],
        [2700, 2400, 4260],
    )

    add_heading(doc, "PHẦN ĐANG TIẾP TỤC THỰC HIỆN", 1, page_break_before=True)
    add_table(
        doc,
        ["Ưu tiên", "Công việc tiếp theo", "Kết quả mong đợi"],
        [
            ["P0", "Hoàn thiện nhận diện object-scope guard cho IDOR qua decorator, dependency và helper liên file", "Giảm false negative mà không tăng cảnh báo thiếu căn cứ"],
            ["P0", "Mở rộng negative case cho mass assignment: DTO allowlist, exclude field, update từng trường", "Đo được false positive và phân biệt pattern an toàn"],
            ["P1", "Review thủ công candidate ngoài ground truth và ghi rõ nguyên nhân TP/FP/FN", "Có dữ liệu hiệu chỉnh rule đáng tin cậy"],
            ["P1", "Chạy pilot LLM thật riêng cho hai nhóm ưu tiên với evidence đã chuẩn hóa", "Đo coverage, mức abstain và lỗi schema/provider"],
            ["P1", "Khóa quality gate trước khi mở rộng phạm vi", "Không dàn trải sang nhóm lỗi logic khác khi hai nhóm hiện tại chưa ổn định"],
        ],
        [900, 5040, 3420],
    )

    add_heading(
        doc,
        "RỦI RO VÀ BIỆN PHÁP KIỂM SOÁT",
        1,
        page_break_before=True,
    )
    add_label_paragraph(
        doc,
        "Thiếu context liên file. ",
        "Guard có thể nằm trong dependency hoặc helper ngoài file. Hướng xử lý là mở rộng "
        "retrieval có giới hạn, lưu lại retrieval-resolution và chuyển MANUAL_REVIEW khi không chứng minh được.",
    )
    add_label_paragraph(
        doc,
        "Ground truth chưa cân bằng. ",
        "Mass assignment hiện thiếu negative case đã review. Báo cáo chỉ công bố số TP/FN, "
        "không suy rộng precision; bước tiếp theo là bổ sung mẫu an toàn đối chứng.",
    )
    add_label_paragraph(
        doc,
        "Tuyên bố quá mức về LLM. ",
        "Kết quả benchmark detector không được gọi là độ chính xác của LLM. Mọi số liệu LLM "
        "sẽ kèm mẫu số, coverage và tỷ lệ NEEDS_REVIEW.",
    )

    add_heading(doc, "KẾT LUẬN TUẦN 4", 1)
    add_body(
        doc,
        "Dự án hiện đã đi qua giai đoạn xây dựng pipeline triage cơ bản và đang ở giai đoạn "
        "phát hiện, thu thập bằng chứng và kiểm chứng lỗi logic. Phạm vi Tuần 4 được giữ ở "
        "hai nhóm: thiếu kiểm soát phân quyền/IDOR và mass assignment. Chức năng tạo candidate "
        "đã hoạt động; công việc gần nhất là cải thiện độ bao phủ, bổ sung negative case, review "
        "candidate ngoài ground truth và đánh giá LLM thật trước khi xem xét mở rộng.",
    )

    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build_report())
