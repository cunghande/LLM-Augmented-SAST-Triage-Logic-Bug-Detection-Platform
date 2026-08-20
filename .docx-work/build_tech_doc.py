from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path(__file__).with_name("SAST_Triage_Tech_Doc_Sprint_1_2.docx")

NAVY = "17365D"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "1F2937"
MUTED = "5F6B7A"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
PALE_BLUE = "F5F8FC"
WHITE = "FFFFFF"
GREEN = "2E6B4F"
AMBER = "8A6400"

CONTENT_DXA = 9360
TABLE_INDENT_DXA = 120


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color="CBD5E1", size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "start", "bottom", "end", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:color"), color)


def set_table_geometry(table, widths_dxa, indent_dxa=TABLE_INDENT_DXA):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            width = widths_dxa[index]
            cell.width = Inches(width / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            set_cell_border(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    if table.rows:
        tr_pr = table.rows[0]._tr.get_or_add_trPr()
        if tr_pr.find(qn("w:tblHeader")) is None:
            header = OxmlElement("w:tblHeader")
            header.set(qn("w:val"), "true")
            tr_pr.append(header)


def set_run_font(run, name="Calibri", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Trang ")
    set_run_font(run, size=8.5, color=MUTED)
    fld_char_1 = OxmlElement("w:fldChar")
    fld_char_1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char_2 = OxmlElement("w:fldChar")
    fld_char_2.set(qn("w:fldCharType"), "end")
    run._r.extend((fld_char_1, instr, fld_char_2))


def configure_numbering(doc):
    numbering = doc.part.numbering_part.element
    abstract_id = 42
    num_id = 42

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    lvl.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "bullet")
    lvl.append(num_fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "•")
    lvl.append(lvl_text)
    suff = OxmlElement("w:suff")
    suff.set(qn("w:val"), "tab")
    lvl.append(suff)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    p_pr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "540")
    ind.set(qn("w:hanging"), "270")
    p_pr.append(ind)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "80")
    spacing.set(qn("w:line"), "300")
    spacing.set(qn("w:lineRule"), "auto")
    p_pr.append(spacing)
    lvl.append(p_pr)
    r_pr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Calibri")
    fonts.set(qn("w:hAnsi"), "Calibri")
    r_pr.append(fonts)
    lvl.append(r_pr)
    abstract.append(lvl)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_num_id = OxmlElement("w:abstractNumId")
    abstract_num_id.set(qn("w:val"), str(abstract_id))
    num.append(abstract_num_id)
    numbering.append(num)
    return num_id


def apply_bullet(paragraph, num_id):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.extend((ilvl, num))


def setup_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    title = styles["Title"]
    title.font.name = "Calibri"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    title.font.size = Pt(25)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(NAVY)
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(4)
    title_p_pr = title._element.get_or_add_pPr()
    title_border = title_p_pr.find(qn("w:pBdr"))
    if title_border is not None:
        title_p_pr.remove(title_border)

    subtitle = styles["Subtitle"]
    subtitle.font.name = "Calibri"
    subtitle._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    subtitle._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    subtitle.font.size = Pt(13)
    subtitle.font.color.rgb = RGBColor.from_string(MUTED)
    subtitle.paragraph_format.space_before = Pt(0)
    subtitle.paragraph_format.space_after = Pt(16)

    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ):
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    code_style = styles.add_style("Code Block", 1)
    code_style.font.name = "Consolas"
    code_style._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
    code_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
    code_style.font.size = Pt(8.4)
    code_style.font.color.rgb = RGBColor.from_string("243447")
    code_style.paragraph_format.left_indent = Inches(0.16)
    code_style.paragraph_format.right_indent = Inches(0.08)
    code_style.paragraph_format.space_before = Pt(5)
    code_style.paragraph_format.space_after = Pt(8)
    code_style.paragraph_format.line_spacing = 1.08


def add_bullet(doc, text, num_id, bold_prefix=None):
    p = doc.add_paragraph()
    apply_bullet(p, num_id)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    if bold_prefix and text.startswith(bold_prefix):
        first = p.add_run(bold_prefix)
        set_run_font(first, bold=True, color=INK)
        rest = p.add_run(text[len(bold_prefix):])
        set_run_font(rest, color=INK)
    else:
        run = p.add_run(text)
        set_run_font(run, color=INK)
    return p


def add_callout(doc, label, text, fill=PALE_BLUE, accent=BLUE):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [CONTENT_DXA])
    cell = table.cell(0, 0)
    shade_cell(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    label_run = p.add_run(f"{label}: ")
    set_run_font(label_run, bold=True, color=accent)
    text_run = p.add_run(text)
    set_run_font(text_run, color=INK)
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(1)


def add_code(doc, text):
    p = doc.add_paragraph(style="Code Block")
    p.paragraph_format.keep_together = True
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), LIGHT_GRAY)
    p_pr.append(shd)
    run = p.add_run(text)
    set_run_font(run, name="Consolas", size=8.4, color="243447")
    return p


def add_comparison_table(doc):
    table = doc.add_table(rows=1, cols=3)
    set_table_geometry(table, [1800, 3780, 3780])
    headers = ("Phần", "Sprint 1", "Sprint 2")
    for index, text in enumerate(headers):
        cell = table.rows[0].cells[index]
        shade_cell(cell, LIGHT_BLUE)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(text)
        set_run_font(run, bold=True, color=NAVY)
    rows = (
        ("Mục tiêu", "Xây pipeline scan và triage an toàn end-to-end.", "Thêm detector logic-bug authorization cho FastAPI."),
        ("Nguồn finding", "Semgrep và Gitleaks.", "Module 2 native AST/index, sau đó nhập pipeline chung."),
        ("Đầu ra", "NormalizedFinding, verdict, policy action, JSON/Markdown.", "Finding CWE-862 cùng verdict/policy/report hiện có."),
        ("Trạng thái", "Luồng sản phẩm chính đã hoạt động.", "Detector chuyên biệt đã triển khai; chưa phải logic-bug tổng quát."),
    )
    for row_data in rows:
        cells = table.add_row().cells
        for index, text in enumerate(row_data):
            p = cells[index].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.12
            run = p.add_run(text)
            set_run_font(run, size=9.7, bold=(index == 0), color=INK)
    set_table_geometry(table, [1800, 3780, 3780])
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)


def build_document():
    doc = Document()
    doc.settings.odd_and_even_pages_header_footer = False
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    setup_styles(doc)
    bullet_num_id = configure_numbering(doc)

    for header_part in (
        section.header,
        section.even_page_header,
        section.first_page_header,
    ):
        header = header_part.paragraphs[0]
        header.clear()
        header.alignment = WD_ALIGN_PARAGRAPH.LEFT
        header.paragraph_format.space_after = Pt(0)
        run = header.add_run("TECHNICAL BRIEF  |  SAST TRIAGE PLATFORM")
        set_run_font(run, size=8.5, color=MUTED, bold=True)
    add_page_number(section.footer.paragraphs[0])

    kicker = doc.add_paragraph()
    kicker.paragraph_format.space_before = Pt(14)
    kicker.paragraph_format.space_after = Pt(4)
    r = kicker.add_run("PROJECT TECH DOC")
    set_run_font(r, size=9.5, color=BLUE, bold=True)

    doc.add_paragraph("LLM-Augmented SAST Triage Platform", style="Title")
    doc.add_paragraph("Tóm tắt kỹ thuật Sprint 1 và Sprint 2", style="Subtitle")

    metadata = (
        ("Phiên bản", "0.1.0"),
        ("Ngày cập nhật", date(2026, 8, 15).strftime("%d/%m/%Y")),
        ("Phạm vi", "Python repository; Semgrep, Gitleaks, FastAPI authorization"),
        ("Trạng thái", "Technical summary - phục vụ trình bày và bàn giao"),
    )
    for label, value in metadata:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        lr = p.add_run(f"{label}: ")
        set_run_font(lr, bold=True, color=NAVY)
        vr = p.add_run(value)
        set_run_font(vr, color=INK)

    doc.add_paragraph()
    add_callout(
        doc,
        "Tóm tắt",
        "Hệ thống nhận một Python repository, chạy scanner, chuẩn hóa finding, lấy ngữ cảnh đã làm sạch, gọi LLM để đề xuất triage, áp dụng policy bảo thủ và xuất báo cáo có thể audit. Sprint 2 bổ sung một detector logic-bug chuyên biệt cho FastAPI thiếu function-level authorization.",
    )

    doc.add_heading("1. Tổng quan hệ thống", level=1)
    p = doc.add_paragraph(
        "Mục tiêu của dự án là hỗ trợ security reviewer phân loại finding mà không để LLM tự động xóa hoặc suppress cảnh báo. Mọi finding đều được giữ trong báo cáo, kể cả khi context thiếu hoặc provider lỗi."
    )
    p.paragraph_format.keep_together = True
    add_code(
        doc,
        "Repository -> Semgrep + Gitleaks -> NormalizedFinding -> Safe Context\n"
        "-> TriageProvider -> PolicyDecision -> report.json + report.md",
    )
    add_comparison_table(doc)

    doc.add_heading("2. Sprint 1 - Nền tảng scan và triage", level=1)
    doc.add_paragraph(
        "Sprint 1 xây dựng pipeline sản phẩm chính. Hai chế độ CLI triage và scan-and-triage khác nhau ở cách lấy scanner artifact nhưng cùng đi qua một pipeline triage duy nhất."
    )

    doc.add_heading("2.1 Các phần đã triển khai", level=2)
    for text, prefix in (
        ("Scanner và parser: chạy Semgrep/Gitleaks, đọc JSON và chuẩn hóa về NormalizedFinding.", "Scanner và parser:"),
        ("Bảo vệ bí mật: Gitleaks redact sớm; raw secret không đi vào model, log, prompt hoặc báo cáo.", "Bảo vệ bí mật:"),
        ("Safe Context: chặn path traversal, absolute/symlink escape, file quá lớn và sanitize credential trong source gần finding.", "Safe Context:"),
        ("Domain model: sáu model public tách biệt, frozen bằng Pydantic v2 và validate ở component boundary.", "Domain model:"),
        ("LLM triage: hỗ trợ fake, Ollama và Qwen3-14B-AWQ trên Modal; output bắt buộc qua schema validation.", "LLM triage:"),
        ("Policy bảo thủ: LLM chỉ đề xuất verdict; policy quyết định LIKELY_TP, FP_CANDIDATE, MANUAL_REVIEW hoặc TRIAGE_ERROR.", "Policy bảo thủ:"),
        ("Báo cáo và audit: xuất report.json, report.md, run-metadata.json, checkpoint và scanner artifacts.", "Báo cáo và audit:"),
    ):
        add_bullet(doc, text, bullet_num_id, prefix)

    doc.add_heading("2.2 Bất biến bảo mật chính", level=2)
    add_bullet(doc, "Raw source chỉ tồn tại trong local scope của SafeContextService.", bullet_num_id)
    add_bullet(doc, "Finding từ repository A không được triage bằng context của repository B.", bullet_num_id)
    add_bullet(doc, "Không lưu raw prompt, raw response, traceback, secret hoặc đường dẫn repository tuyệt đối trong report.", bullet_num_id)
    add_bullet(doc, "Lỗi trên một finding không được làm mất các finding còn lại.", bullet_num_id)

    doc.add_heading("2.3 Chạy Sprint 1", level=2)
    doc.add_paragraph("Điều kiện: .venv đã active và PowerShell đang ở D:\\Code\\target-repos.")
    add_code(doc, "python -m sast_triage TEN_REPO")
    doc.add_paragraph(
        "Lệnh rút gọn dùng provider real trên Modal và tự tạo output có timestamp trong D:\\Code\\triage-runs. Dùng provider fake khi chỉ cần kiểm tra scanner, pipeline và report offline."
    )

    doc.add_heading("3. Sprint 2 - FastAPI authorization logic detector", level=1)
    add_callout(
        doc,
        "Phạm vi thực tế",
        "Sprint 2 hiện phát hiện một lớp logic-bug: endpoint FastAPI nhạy cảm thiếu function-level authorization (CWE-862). Đây chưa phải detector tổng quát cho năm lớp logic bug.",
        fill="FFF8E6",
        accent=AMBER,
    )

    doc.add_heading("3.1 Cách detector hoạt động", level=2)
    for text in (
        "Dùng Python AST và bounded repository index để nhận diện FastAPI, APIRouter, route decorator, Depends, Security và include_router.",
        "Đánh giá route có nhạy cảm hay không dựa trên method, resource identifier, route semantics và control đã biết.",
        "Phân biệt authentication với authorization; chỉ có get_current_user không đồng nghĩa người dùng có quyền thao tác tài nguyên.",
        "Phân loại nội bộ thành PROTECTED, MISSING_AUTHZ_CANDIDATE hoặc UNKNOWN.",
        "Candidate hoàn chỉnh trở thành finding module2.fastapi.missing-function-authz và đi qua Safe Context, LLM, policy, reporter của Sprint 1.",
    ):
        add_bullet(doc, text, bullet_num_id)

    doc.add_heading("3.2 Quy tắc giảm false positive", level=2)
    add_bullet(doc, "Route public như health, login, docs và metrics không bị báo nếu không có dấu hiệu nhạy cảm khác.", bullet_num_id)
    add_bullet(doc, "Router-level dependency và include_router dependency được kế thừa khi đánh giá authorization.", bullet_num_id)
    add_bullet(doc, "Dynamic middleware hoặc context không đủ được chuyển thành UNKNOWN thay vì tự kết luận thiếu authorization.", bullet_num_id)
    add_bullet(doc, "Không tự động suppress finding; reviewer vẫn quyết định cuối cùng.", bullet_num_id)

    doc.add_heading("3.3 Chạy Sprint 2", level=2)
    doc.add_paragraph(
        "Điều kiện: .venv đã active, Semgrep/Gitleaks có trong PATH và PowerShell đang ở D:\\Code\\target-repos. Đổi cả hai chỗ TEN_REPO thành tên thư mục repository FastAPI."
    )
    add_code(
        doc,
        "python -m sast_triage scan-and-triage --repository .\\TEN_REPO --output-dir "
        "..\\triage-runs\\TEN_REPO-sprint2 --analysis-backend native --module2 "
        "--logic-detector fastapi-missing-authz --provider fake",
    )
    doc.add_paragraph("Dùng LLM thật trên Modal bằng cách đổi --provider fake thành --provider real.")

    doc.add_heading("4. Kết quả và trạng thái hiện tại", level=1)
    doc.add_heading("4.1 Qualification detector Sprint 2", level=2)
    p = doc.add_paragraph(
        "Bộ qualification hiện tại có 80 case (35 natural RealVuln và 45 controlled mutation), cân bằng 40 positive/40 negative. Kết quả ghi nhận TP=40, FP=0, FN=0, TN=39 và UNKNOWN=1; held-out có 38 case với TP=20, FP=0, FN=0, TN=18."
    )
    p.paragraph_format.keep_together = True
    add_callout(
        doc,
        "Cảnh báo diễn giải",
        "Các chỉ số trên chỉ đúng với dataset qualification hiện tại. Governance gate vẫn chờ independent two-reviewer audit; không dùng kết quả fixture hoặc mutation để tuyên bố chất lượng tổng quát ngoài dataset.",
        fill=LIGHT_GRAY,
        accent=MUTED,
    )

    doc.add_heading("4.2 Kiểm thử và hồi quy", level=2)
    add_bullet(doc, "Ba run detector deterministic cho kết quả giống nhau từng byte.", bullet_num_id)
    add_bullet(doc, "Paired offline regression: 20 repository x 3 lần x 2 nhánh, tổng 120 run; không ghi nhận pipeline error trong báo cáo qualification.", bullet_num_id)
    add_bullet(doc, "Security suite: 175 passed, 2 skipped; full suite: 488 passed, 2 skipped theo báo cáo hiện tại.", bullet_num_id)

    doc.add_heading("4.3 Vấn đề vận hành đã quan sát", level=2)
    doc.add_paragraph(
        "Run gần nhất trên Sequential-Recommender tìm được 13 finding (1 Semgrep và 12 Module 2), nhưng Modal client ngắt kết nối trong lúc vLLM/FlashInfer warm-up. Vì vậy 13 finding được giữ lại với TRIAGE_ERROR; scanner và detector vẫn hoàn tất, còn LLM chưa tạo verdict."
    )

    doc.add_heading("5. Đầu ra và cách đọc", level=1)
    for text, prefix in (
        ("report.md: báo cáo chính để reviewer và mentor đọc.", "report.md:"),
        ("report.json: dữ liệu finding, decision, policy và lỗi xử lý có cấu trúc.", "report.json:"),
        ("run-metadata.json: commit, scanner version/config, provider, model và repository binding.", "run-metadata.json:"),
        ("scanner-artifacts/: JSON gốc từ Semgrep/Gitleaks; có thể nhạy cảm và không nên gửi ra ngoài tùy tiện.", "scanner-artifacts/:")
    ):
        add_bullet(doc, text, bullet_num_id, prefix)

    limits_heading = doc.add_heading("6. Giới hạn và bước tiếp theo", level=1)
    limits_heading.paragraph_format.page_break_before = True
    add_bullet(doc, "Chưa phát hiện tổng quát IDOR/BOLA, mass assignment, race condition/TOCTOU, cross-service authorization hoặc business workflow bypass.", bullet_num_id)
    add_bullet(doc, "Chưa có auto-fix, auto-suppression, risk scoring đã hiệu chỉnh, dashboard hoặc CI/CD gate.", bullet_num_id)
    add_bullet(doc, "Cần hoàn tất independent reviewer audit và real-LLM regression trước khi công bố quality claim rộng hơn.", bullet_num_id)
    add_bullet(doc, "Nên cải thiện Modal provider theo hướng fail-fast, retry có giới hạn và giữ checkpoint khi remote session bị ngắt.", bullet_num_id)

    doc.add_heading("7. Tài liệu nguồn nội bộ", level=1)
    for text in (
        "docs/sprint-1-scope.md - phạm vi và iteration Sprint 1.",
        "docs/sprint-2-scope.md - detector FastAPI authorization và qualification status.",
        "docs/architecture.md - pipeline, component boundary và repository-artifact binding.",
        "docs/cli-usage.md - cài đặt, câu lệnh chạy, output và troubleshooting.",
        "docs/security-policy.md - quy tắc xử lý source, secret, prompt và report.",
    ):
        add_bullet(doc, text, bullet_num_id)

    core = doc.core_properties
    core.title = "LLM-Augmented SAST Triage Platform - Tech Doc Sprint 1 và Sprint 2"
    core.subject = "Tóm tắt kỹ thuật các phần đã triển khai"
    core.author = "SAST Triage Project"
    core.keywords = "SAST, Semgrep, Gitleaks, LLM, FastAPI, authorization, Sprint 1, Sprint 2"
    core.comments = "Technical summary generated from repository documentation."

    doc.settings.odd_and_even_pages_header_footer = False
    section.different_first_page_header_footer = False

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
